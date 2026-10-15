"""Lecture du référentiel des modes d'intervention (fichier Excel ART enrichi)."""
from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl

URL_RE = re.compile(r"https?://[^\s;,)]+")


@dataclass
class ModeIntervention:
    id: str
    famille: str
    nom: str
    exemples: list[str] = field(default_factory=list)  # dispositifs cités entre parenthèses
    dispositifs_evalues: list[str] = field(default_factory=list)

    @property
    def libelle_court(self) -> str:
        """Nom sans la parenthèse finale, ex. 'Prêt bonifié à taux zéro'."""
        return re.split(r"\s*\(", self.nom, maxsplit=1)[0].strip()

    def termes_recherche(self, maximum: int = 4) -> list[str]:
        """Termes concrets pour la recherche web : libellé + dispositifs nommés."""
        termes = [self.libelle_court]
        for t in self.dispositifs_evalues + self.exemples:
            t = re.sub(r"\(.*?\)", "", t)
            t = t.replace("(", "").replace(")", "").strip(" .:…")
            sigle = t.isupper() and len(t) >= 3
            if not (sigle or 6 <= len(t) <= 60):
                continue  # trop court pour une requête utile ("Air", "masse"...)
            if t.lower() not in (x.lower() for x in termes):
                termes.append(t)
        return termes[:maximum]

    def pour_prompt(self) -> str:
        return f"{self.id} | {self.famille} | {self.nom}"


def _exemples(nom: str) -> list[str]:
    m = re.search(r"\((.*)\)", nom)
    if not m:
        return []
    contenu = m.group(1)
    contenu = re.sub(r"\bex\.\s*", "", contenu)
    morceaux = re.split(r",|;|\bpuis\b|/|\bet\b", contenu)
    return [x.strip(" .…") for x in morceaux if x.strip(" .…")]


def charger_referentiel(chemin: str) -> tuple[list[ModeIntervention], list[str]]:
    """Retourne (liste des MI, liste d'URL sources citées dans le fichier)."""
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    mis: list[ModeIntervention] = []
    for ligne in ws.iter_rows(min_row=3, values_only=True):
        famille, nom = (ligne[0] or ""), (ligne[1] or "")
        famille, nom = str(famille).strip(), str(nom).strip()
        if not famille or not nom:
            continue  # lignes de section ou de légende
        mis.append(ModeIntervention(id=f"MI{len(mis) + 1:02d}", famille=famille,
                                    nom=nom, exemples=_exemples(nom)))

    # Feuille des évaluations : dispositifs évalués + URL
    urls: list[str] = []
    for ws in wb.worksheets[1:]:
        entetes = None
        for ligne in ws.iter_rows(values_only=True):
            cellules = [str(c) if c is not None else "" for c in ligne]
            if entetes is None and any(c in ("Mode d'intervention", "Document") for c in cellules):
                entetes = cellules
                continue
            for c in cellules:
                urls += URL_RE.findall(c)
            if entetes and "Dispositif évalué" in entetes and "Mode d'intervention" in entetes:
                mi_txt = cellules[entetes.index("Mode d'intervention")].lower()
                disp = cellules[entetes.index("Dispositif évalué")]
                for mi in mis:
                    if mi.libelle_court.lower()[:25] in mi_txt or mi_txt[:25] in mi.nom.lower():
                        mi.dispositifs_evalues += [d.strip() for d in re.split(r",|;", disp) if d.strip()]
                        break
    # dédoublonnage en gardant l'ordre
    vus, urls_uniques = set(), []
    for u in urls:
        u = u.rstrip(".")
        if u not in vus:
            vus.add(u)
            urls_uniques.append(u)
    return mis, urls_uniques
